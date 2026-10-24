"""
MetaShield - Realistic Metadata Test Sample Generator
Generates test files with embedded GPS, device fingerprints, author PII, revisions, and copyright tags.
"""

import os
from pathlib import Path
from PIL import Image, ExifTags
import pypdf
import docx

def generate_samples(output_dir: str = "engine/test_samples"):
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    # 1. Realistic JPEG with GPS, Camera Serial, and Copyright
    img_path = out / "sample_field_photo.jpg"
    img = Image.new("RGB", (320, 240), color=(15, 23, 42))
    
    exif = Image.Exif()
    # Make & Model
    exif[0x010F] = "Nikon Corporation"  # Make
    exif[0x0110] = "NIKON Z 9"          # Model
    exif[0x0131] = "Adobe Photoshop 2026"  # Software
    exif[0x013B] = "Dr. Alexander Vance <avance@mru.edu.in>"  # Artist
    exif[0x8298] = "Copyright (C) 2026 MetaShield Lab, Malla Reddy University"  # Copyright
    exif[0x9003] = "2026:09:23 11:30:00"  # DateTimeOriginal
    exif[0xA431] = "SN-9482104-Z9"       # CameraSerialNumber

    # GPS IFD: 37°46'29.8"N 122°25'09.6"W (San Francisco)
    gps_ifd = exif.get_ifd(ExifTags.IFD.GPSInfo)
    gps_ifd[1] = "N"
    gps_ifd[2] = (37.0, 46.0, 29.8)
    gps_ifd[3] = "W"
    gps_ifd[4] = (122.0, 25.0, 9.6)
    gps_ifd[6] = 15.0 # Altitude 15m

    img.save(str(img_path), "JPEG", exif=exif)
    print(f"Generated: {img_path}")

    # 2. Realistic DOCX with Authors, Comments, and Revision Info
    docx_path = out / "sample_executive_brief.docx"
    doc = docx.Document()
    doc.add_heading("MetaShield Operational Security Memo", level=1)
    doc.add_paragraph(
        "This confidential briefing discusses privacy minimization procedures under GDPR Article 5(1)(c) "
        "and data sanitization frameworks developed at Malla Reddy University by Charan Konda and J. Naga Phani."
    )
    doc.add_paragraph("All sensitive telemetry and GPS markers must be purged prior to public publication.")
    
    cp = doc.core_properties
    cp.author = "Charan Konda (2511CS040113)"
    cp.last_modified_by = "J. Naga Phani (2511CS040132)"
    cp.title = "MetaShield Security Briefing"
    cp.subject = "Compliance & Privacy Architecture"
    cp.comments = "INTERNAL DRAFT ONLY - Contains confidential architectural blueprints."
    cp.revision = 14
    
    doc.save(str(docx_path))
    print(f"Generated: {docx_path}")

    # 3. Realistic PDF with PDFInfo and metadata
    pdf_path = out / "sample_compliance_memo.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_metadata({
        "/Title": "MetaShield Compliance Assessment",
        "/Author": "Investigator Vance (ID: 8841-SEC)",
        "/Subject": "HIPAA & GDPR Metadata Audit",
        "/Creator": "Acrobat PDF Maker 2026",
        "/Producer": "Malla Reddy University Document System",
        "/Copyright": "Copyright (C) 2026 Malla Reddy University"
    })
    with open(str(pdf_path), "wb") as f_pdf:
        writer.write(f_pdf)
    print(f"Generated: {pdf_path}")

if __name__ == "__main__":
    generate_samples()
